### rrrr
import json, re, subprocess, sys, tempfile
from pathlib import Path
import streamlit as st

BASE = Path(__file__).parent
TEMPLATE = BASE / 'assets' / 'Candidate_Template.docx'
BUILDER = BASE / 'scripts' / 'build_cv.py'

st.set_page_config(page_title='Ford / Magnit CV Generator', page_icon='📄', layout='wide')
st.title('Ford / Magnit CV Generator')
st.caption('Upload a candidate resume. The app extracts resume content, validates the submission fields, and generates the editable Ford/Magnit Word CV.')

with st.sidebar:
    st.header('Rules')
    st.write('• Resume-only extraction. No invented facts.')
    st.write('• Missing Ford/certification/score fields → NA')
    st.write('• Missing scheduling fields → TBC')
    st.write('• Exactly 3 top skills')
    st.write('• Employment: most recent first')
    st.write('• Appended resume excludes contact details, links, photo and logos')
    st.write('• Candidate photo is optional and must be supplied separately')

resume_file = st.file_uploader('Candidate resume', type=['pdf','docx','txt'])
photo_file = st.file_uploader('Optional candidate photo for the page-1 photo box', type=['jpg','jpeg','png'])

st.info('This version keeps the original build_cv.py and Candidate_Template.docx from the uploaded skill package. Resume extraction is deliberately conservative. Review the generated fields before sending to a client.')

if resume_file:
    suffix = Path(resume_file.name).suffix.lower()
    tmpdir = Path(tempfile.mkdtemp(prefix='ford_cv_'))
    resume_path = tmpdir / ('resume' + suffix)
    resume_path.write_bytes(resume_file.getvalue())

    try:
        if suffix == '.pdf':
            from pypdf import PdfReader
            reader = PdfReader(str(resume_path))
            text = '\n'.join((p.extract_text() or '') for p in reader.pages)
        elif suffix == '.docx':
            from docx import Document
            doc = Document(str(resume_path))
            text = '\n'.join(p.text for p in doc.paragraphs)
            for table in doc.tables:
                for row in table.rows:
                    text += '\n' + ' | '.join(cell.text for cell in row.cells)
        else:
            text = resume_path.read_text(errors='ignore')
    except Exception as e:
        st.error(f'Could not read the resume: {e}')
        st.stop()

    if not text.strip():
        st.error('No readable text was extracted from the resume. For scanned PDFs, OCR the resume first.')
        st.stop()

    st.subheader('Extracted resume text')
    st.text_area('Review source text', text, height=280)

    st.subheader('Candidate data')
    st.caption('Paste or edit candidate.json below. This preserves the skill package schema and gives you control over any ambiguous fields.')
    example = json.loads((BASE / 'assets' / 'example_candidate.json').read_text())
    example_json = json.dumps(example, indent=2, ensure_ascii=False)
    candidate_json = st.text_area('candidate.json', example_json, height=500)

    col1, col2 = st.columns(2)
    with col1:
        st.write('**Validation checklist**')
        st.write('✓ Company names must match the resume wording exactly')
        st.write('✓ Do not invent dates, experience, certifications or notice period')
        st.write('✓ Keep scheduling values short: TBC, Immediate, 7 Days, 15 Days, etc.')
    with col2:
        st.write('**Output**')
        st.write('Fixed Ford/Magnit template first, followed by editable resume text when the `resume` block is included.')

    if st.button('Generate Ford / Magnit CV', type='primary'):
        try:
            data = json.loads(candidate_json)
        except json.JSONDecodeError as e:
            st.error(f'Invalid JSON: {e}')
            st.stop()

        if len(data.get('top_skills', [])) != 3:
            st.error('top_skills must contain exactly 3 skills.')
            st.stop()

        for k in ['notice_period','internal_external','interview_availability','start_availability']:
            if len(str(data.get(k, 'TBC'))) > 24:
                st.warning(f'{k} is long. The template has limited right-side space and may wrap.')

        data_path = tmpdir / 'candidate.json'
        data_path.write_text(json.dumps(data, indent=2, ensure_ascii=False))
        output_path = tmpdir / f"{data.get('first_name','Candidate')}_{data.get('last_name','')}_Ford_Magnit_CV.docx"
        output_path = Path(re.sub(r'[^A-Za-z0-9_.-]+', '_', output_path.name))
        output_path = tmpdir / output_path.name

        cmd = [sys.executable, str(BUILDER), '--data', str(data_path), '--output', str(output_path), '--template', str(TEMPLATE)]
        photo_path = None
        if photo_file:
            photo_path = tmpdir / ('photo' + Path(photo_file.name).suffix.lower())
            photo_path.write_bytes(photo_file.getvalue())
            cmd += ['--photo', str(photo_path)]

        result = subprocess.run(cmd, capture_output=True, text=True, cwd=str(BASE))
        if result.returncode != 0:
            st.error('CV generation failed.')
            st.code(result.stdout + '\n' + result.stderr)
            st.stop()

        st.success('CV generated successfully.')
        if result.stdout:
            st.code(result.stdout)
        st.download_button('Download Ford / Magnit CV', output_path.read_bytes(), file_name=output_path.name, mime='application/vnd.openxmlformats-officedocument.wordprocessingml.document')

        missing = []
        for k in ['overall_it_experience','experience_core_skill','other_qualifications','hackerrank_score','certifications','education_master','education_master_duration']:
            if str(data.get(k, 'NA')).strip().upper() == 'NA':
                missing.append(k)
        for k in ['notice_period','internal_external','interview_availability','start_availability']:
            if str(data.get(k, 'TBC')).strip().upper() == 'TBC':
                missing.append(k)
        if missing:
            st.warning('Fields still marked NA/TBC: ' + ', '.join(missing))
